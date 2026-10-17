"""Spreadsheet loader: openpyxl (xlsx) + stdlib csv -> text table with cell refs."""

from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import load_workbook

from ..models import DocumentText, Page


def load(path: str | Path, doc_id: str | None = None) -> DocumentText:
    p = Path(path)
    doc_id = doc_id or p.stem
    if p.suffix.lower() == ".csv":
        return _load_csv(p, doc_id)
    return _load_xlsx(p, doc_id)


def _load_xlsx(path: Path, doc_id: str) -> DocumentText:
    wb = load_workbook(path, data_only=True)
    pages: list[Page] = []
    for sheet in wb.worksheets:
        lines = []
        for row in sheet.iter_rows():
            cells = []
            for cell in row:
                if cell.value is not None:
                    cells.append(f"{cell.coordinate}={cell.value}")
            if cells:
                lines.append(" | ".join(cells))
        pages.append(Page(ref=sheet.title, text="\n".join(lines), ocr_confidence=None))
    wb.close()
    return DocumentText(doc_id=doc_id, doc_type="spreadsheet", pages=pages, source_path=str(path.resolve()))


def _load_csv(path: Path, doc_id: str) -> DocumentText:
    lines = []
    with open(path, newline="", encoding="utf-8") as f:
        reader = csv.reader(f)
        for r, row in enumerate(reader, start=1):
            cells = [f"{_col(c_idx)}{r}={val}" for c_idx, val in enumerate(row) if val != ""]
            if cells:
                lines.append(" | ".join(cells))
    page = Page(ref="Sheet1", text="\n".join(lines), ocr_confidence=None)
    return DocumentText(doc_id=doc_id, doc_type="spreadsheet", pages=[page], source_path=str(path.resolve()))


def _col(idx: int) -> str:
    """0-based index -> spreadsheet column letter(s) (A, B, ... Z, AA...)."""
    s = ""
    idx += 1
    while idx:
        idx, rem = divmod(idx - 1, 26)
        s = chr(65 + rem) + s
    return s


def marked_text(doc: DocumentText) -> str:
    out = []
    for pg in doc.pages:
        out.append(f"[{pg.ref}]")
        out.append(pg.text)
    return "\n".join(out)
