"""Shared fixtures: tiny PDFs (text + scanned), xlsx, eml."""

from __future__ import annotations

import io
from pathlib import Path

import pytest


@pytest.fixture(autouse=True)
def _isolated_review_db(tmp_path, monkeypatch):
    """Never let tests write into the real data/review.db."""
    monkeypatch.setattr("invoicelens.review_queue.DEFAULT_DB", tmp_path / "review.db")


@pytest.fixture(scope="session")
def fixtures_dir(tmp_path_factory) -> Path:
    d = tmp_path_factory.mktemp("fixtures")
    _make_text_pdf(d / "text.pdf")
    _make_scanned_pdf(d / "scanned.pdf")
    _make_xlsx(d / "sheet.xlsx")
    _make_eml(d / "mail.eml")
    return d


def _make_text_pdf(path: Path):
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(str(path), pagesize=letter)
    c.drawString(72, 720, "ACME Supplies Ltd")
    c.drawString(72, 700, "INVOICE INV-1001")
    c.drawString(72, 680, "Date: 2026-03-04")
    c.drawString(72, 660, "Currency: USD")
    c.drawString(72, 640, "Widget x2 @ 10.00 = 20.00")
    c.drawString(72, 620, "TOTAL: 22.00")
    c.save()


def _make_scanned_pdf(path: Path):
    """Image-only PDF (no extractable text) -> router calls it scanned_pdf."""
    import fitz
    from PIL import Image

    img = Image.new("RGB", (400, 300), (255, 255, 255))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    doc = fitz.open()
    page = doc.new_page(width=400, height=300)
    page.insert_image(page.rect, stream=buf.getvalue())
    doc.save(str(path))
    doc.close()


def _make_xlsx(path: Path):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Invoice"
    ws["A1"] = "ACME Supplies Ltd"
    ws["A2"] = "INVOICE"
    ws["B3"] = "INV-1001"
    ws["B6"] = "USD"
    ws["A9"] = "Description"
    ws["A10"] = "Widget"
    ws["B10"] = 2
    ws["C10"] = 10.0
    ws["D10"] = 20.0
    ws["D13"] = 22.0
    wb.save(str(path))


def _make_eml(path: Path):
    import email.message

    msg = email.message.EmailMessage()
    msg["From"] = "a@example.com"
    msg["To"] = "b@example.com"
    msg["Subject"] = "Invoice INV-1001"
    msg.set_content(
        "Supplier: ACME Supplies Ltd\nDocument No: INV-1001\nDate: 2026-03-04\nCurrency: USD\nTOTAL: 22.00"
    )
    path.write_bytes(bytes(msg))
